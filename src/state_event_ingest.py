"""Ingester for state-mandated adverse-event registries (currently MA SREs).

Massachusetts Department of Public Health publishes annual XLSX workbooks
listing every Serious Reportable Event (NQF SRE list) at every licensed
acute-care hospital, non-acute hospital, and ambulatory surgical center.
This is the only public US source we've verified that ties a *named hospital*
to a *device-relevant adverse event* — public MAUDE has no hospital identifier.

The workbook URL pattern (verified live April 2026):
    https://www.mass.gov/doc/calendar-year-{YEAR}-serious-reportable-events
       -in-massachusetts-{VARIANT}/download
where YEAR ∈ 2015..2022 and VARIANT ∈
    acute-care-hospitals | non-acute-care-hospitals | ambulatory-surgical-centers

Layouts diverge slightly:
  acute: 2-row header (event-group banner + per-event column), 28 event types.
  non_acute / asc: 1-row header, fewer event types.

CCN matching: hospital_name from MA → cms_hospitals.facility_name where
state='MA'. Acute-care hospitals reliably resolve; non-acute and ASCs mostly
do not (they aren't in cms_hospitals).
"""
from __future__ import annotations

import logging
import os
import re
from dataclasses import dataclass
from typing import Iterable, Optional

import openpyxl
import requests

from . import db

log = logging.getLogger(__name__)

DOWNLOAD_DIR = "downloads/state_events/ma"
URL_TEMPLATE = (
    "https://www.mass.gov/doc/calendar-year-{year}-serious-reportable-events"
    "-in-massachusetts-{variant}/download"
)

VARIANTS = {
    "acute":     ("acute-care-hospitals",         "acute"),
    "non_acute": ("non-acute-care-hospitals",     "non_acute"),
    "asc":       ("ambulatory-surgical-centers",  "asc"),
}

# Years discovered live (2026-04-28) via HEAD probe. ASC didn't begin
# reporting until 2017. Update this list when MA publishes 2023+.
YEARS_BY_VARIANT = {
    "acute":     range(2015, 2023),
    "non_acute": range(2015, 2023),
    "asc":       range(2017, 2023),
}

# Event-type keywords that flag a row as device-related. Conservative: only
# flag categories where a medical device is the actual or proximate cause.
# We deliberately include "Retained foreign object" — NJ's parallel registry
# explicitly classifies "unretrieved device fragment" under PSI 05.
DEVICE_KEYWORDS = (
    "device",
    "retained foreign object",
    "metallic object in mri",
    "wrong or contaminated o2",
    "restraint or bedrail",
    "electrical shock",
    "burn w/death",
    "burn with death",
    "intravascular air embolism",
)


# ------------------------------------------------------------------
# Hospital-name normalization + CCN matching
# ------------------------------------------------------------------

_LICENSE_ID_RE = re.compile(r"\(([A-Za-z0-9]+)\)\s*$")
_PUNCT_RE = re.compile(r"[^a-z0-9 ]+")
_WS_RE = re.compile(r"\s+")

# Common name expansions/strippings. Apply before tokenization.
_REPLACEMENTS = (
    (re.compile(r"\bmass\b", re.I),    "massachusetts"),
    (re.compile(r"\bmed\b", re.I),     "medical"),
    (re.compile(r"\bctr\b", re.I),     "center"),
    (re.compile(r"\bhosp\b", re.I),    "hospital"),
    (re.compile(r"\bisr\b", re.I),     "israel"),
    (re.compile(r"\bdeacnss\b", re.I), "deaconess"),
    (re.compile(r"\bsts?\b", re.I),    "saint"),
    (re.compile(r"\b&\b"),             "and"),
)

# Suffixes/noise tokens to drop after replacement.
_NOISE_TOKENS = {
    "inc", "incorporated", "llc", "corp", "corporation", "company", "co",
    "ltd", "lp", "the", "a", "of", "campus", "campu", "campusu",
    "addison", "gilbert",  # campus suffixes are too noisy on their own
}


def _strip_license_id(name: str) -> tuple[str, Optional[str]]:
    """Pull the trailing license id like '(2168)' off a name. Returns (clean_name, id)."""
    m = _LICENSE_ID_RE.search(name)
    if not m:
        return name.strip(), None
    return name[:m.start()].strip(), m.group(1)


def _normalize(name: str) -> str:
    s = name.lower()
    for pat, repl in _REPLACEMENTS:
        s = pat.sub(repl, s)
    s = _PUNCT_RE.sub(" ", s)
    s = _WS_RE.sub(" ", s).strip()
    return s


def _tokens(name: str) -> set[str]:
    return {t for t in _normalize(name).split() if t and t not in _NOISE_TOKENS}


def _build_ccn_index(state: str) -> list[tuple[str, str, set[str]]]:
    """Return [(ccn, normalized_name, token_set), ...] for one state."""
    rows: list[tuple[str, str, set[str]]] = []
    with db.connect() as conn:
        res = conn.query(
            "SELECT facility_id, facility_name FROM cms_hospitals FINAL "
            f"WHERE state = '{state}'"
        )
        for ccn, fname in res.result_rows:
            if not fname:
                continue
            rows.append((str(ccn), _normalize(str(fname)), _tokens(str(fname))))
    return rows


def _match_ccn(state_name: str, index: list[tuple[str, str, set[str]]]) -> tuple[Optional[str], str]:
    """Return (ccn, confidence). confidence ∈ {'exact','fuzzy','none'}."""
    norm = _normalize(state_name)
    toks = _tokens(state_name)
    if not toks:
        return None, "none"
    # 1. Exact normalized hit.
    for ccn, n, _ in index:
        if n == norm:
            return ccn, "exact"
    # 2. Substring containment either way (handles "Mass General Hospital"
    #    vs "Massachusetts General Hospital" once 'mass'→'massachusetts').
    for ccn, n, _ in index:
        if norm and n and (norm in n or n in norm):
            return ccn, "fuzzy"
    # 3. Token-overlap ≥ 60% AND ≥ 2 distinguishing tokens shared.
    best_ccn, best_overlap = None, 0.0
    for ccn, _, t in index:
        if not t:
            continue
        shared = toks & t
        if len(shared) < 2:
            continue
        overlap = len(shared) / max(len(toks), len(t))
        if overlap > best_overlap:
            best_overlap, best_ccn = overlap, ccn
    if best_ccn and best_overlap >= 0.6:
        return best_ccn, "fuzzy"
    return None, "none"


# ------------------------------------------------------------------
# XLSX parsing
# ------------------------------------------------------------------

@dataclass
class FactRow:
    state: str
    report_year: int
    facility_type: str
    hospital_name: str
    state_license_id: Optional[str]
    event_category: str
    event_type: str
    is_device_related: bool
    event_count: int
    source_url: str
    source_doc_label: str


def _is_device(event_type: str, event_category: str) -> bool:
    haystack = f"{event_type} {event_category}".lower()
    if event_category.strip().lower() == "product or device events":
        return True
    return any(k in haystack for k in DEVICE_KEYWORDS)


def _parse_workbook(path: str, *, year: int, facility_type: str,
                    source_url: str, source_doc_label: str) -> list[FactRow]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    # Acute-care has a 2-row header (category banner + event sub-header).
    # Non-acute and ASC have a 1-row header.
    if facility_type == "acute":
        category_row = rows[0]
        event_row = rows[1]
        data_rows = rows[2:]
        # Forward-fill the merged category banner.
        last_cat = ""
        categories = []
        for v in category_row:
            if v:
                last_cat = str(v).strip()
            categories.append(last_cat)
    else:
        event_row = rows[0]
        data_rows = rows[1:]
        categories = ["" for _ in event_row]

    out: list[FactRow] = []
    for row in data_rows:
        name_cell = row[0]
        if not name_cell:
            continue
        name = str(name_cell).strip()
        if not name or name.lower() in ("grand total", "total"):
            continue
        clean_name, license_id = _strip_license_id(name)
        for col_idx in range(1, len(event_row)):
            ev = event_row[col_idx]
            if not ev:
                continue
            ev_str = str(ev).strip()
            if ev_str.lower() == "grand total":
                continue
            cat = categories[col_idx] if col_idx < len(categories) else ""
            cell = row[col_idx] if col_idx < len(row) else None
            try:
                count = int(cell) if cell not in (None, "", " ") else 0
            except (TypeError, ValueError):
                count = 0
            if count == 0:
                continue
            out.append(FactRow(
                state="MA",
                report_year=year,
                facility_type=facility_type,
                hospital_name=clean_name,
                state_license_id=license_id,
                event_category=cat or ev_str,  # fall back to event_type if no group
                event_type=ev_str,
                is_device_related=_is_device(ev_str, cat or ""),
                event_count=count,
                source_url=source_url,
                source_doc_label=source_doc_label,
            ))
    return out


# ------------------------------------------------------------------
# Download + ingest pipeline
# ------------------------------------------------------------------

def _download(year: int, variant_slug: str, facility_type: str) -> Optional[str]:
    os.makedirs(DOWNLOAD_DIR, exist_ok=True)
    path = os.path.join(DOWNLOAD_DIR, f"{year}_{facility_type}.xlsx")
    if os.path.exists(path) and os.path.getsize(path) > 1024:
        return path
    url = URL_TEMPLATE.format(year=year, variant=variant_slug)
    try:
        r = requests.get(url, timeout=60)
    except requests.RequestException as e:
        log.warning("MA SRE %s/%s: download failed: %s", year, facility_type, e)
        return None
    if r.status_code != 200 or len(r.content) < 1024:
        log.info("MA SRE %s/%s: not published (HTTP %s, %d bytes)",
                 year, facility_type, r.status_code, len(r.content))
        return None
    with open(path, "wb") as f:
        f.write(r.content)
    return path


def _insert(rows: Iterable[FactRow], ccn_index: list[tuple[str, str, set[str]]]) -> int:
    batch: list[tuple] = []
    for r in rows:
        ccn, conf = _match_ccn(r.hospital_name, ccn_index)
        batch.append((
            r.state, r.report_year, r.facility_type,
            r.hospital_name, r.state_license_id,
            ccn, conf,
            r.event_category, r.event_type, 1 if r.is_device_related else 0,
            r.event_count, r.source_url, r.source_doc_label,
        ))
    if not batch:
        return 0
    with db.connect() as conn:
        conn.insert(
            "state_adverse_events",
            batch,
            column_names=[
                "state", "report_year", "facility_type",
                "hospital_name", "state_license_id",
                "ccn_match", "ccn_match_confidence",
                "event_category", "event_type", "is_device_related",
                "event_count", "source_url", "source_doc_label",
            ],
        )
    return len(batch)


def ingest_ma(years: Optional[list[int]] = None,
              variants: Optional[list[str]] = None) -> dict:
    """Run the full MA SRE ingest. Returns per-variant row counts."""
    db.apply_schema()
    ccn_index = _build_ccn_index("MA")
    log.info("Built CCN index for MA: %d facilities", len(ccn_index))

    selected_variants = variants or list(VARIANTS.keys())
    summary: dict = {}
    for facility_type in selected_variants:
        if facility_type not in VARIANTS:
            log.warning("unknown variant: %s", facility_type)
            continue
        url_slug, _ = VARIANTS[facility_type]
        year_range = years or list(YEARS_BY_VARIANT[facility_type])
        v_rows = 0
        v_match = 0
        for year in year_range:
            path = _download(year, url_slug, facility_type)
            if not path:
                continue
            url = URL_TEMPLATE.format(year=year, variant=url_slug)
            label = f"CY {year} {facility_type.replace('_', '-')} ({url_slug})"
            try:
                facts = _parse_workbook(
                    path, year=year, facility_type=facility_type,
                    source_url=url, source_doc_label=label,
                )
            except Exception as e:
                log.exception("MA SRE %s/%s parse failed: %s", year, facility_type, e)
                continue
            if not facts:
                continue
            inserted = _insert(facts, ccn_index)
            matched = sum(
                1 for f in facts
                if _match_ccn(f.hospital_name, ccn_index)[0] is not None
            )
            v_rows += inserted
            v_match += matched
            log.info("MA SRE %s/%s: %d fact rows (%d CCN-matched)",
                     year, facility_type, inserted, matched)
        summary[facility_type] = {"rows": v_rows, "ccn_matched": v_match}
    return summary


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(message)s")
    s = ingest_ma()
    for k, v in s.items():
        print(f"  {k}: {v['rows']} rows · {v['ccn_matched']} CCN-matched")
